import os
import io
from typing import List
from fastapi import FastAPI, Depends, HTTPException, Request, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, FileResponse, Response
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session

import models
import schemas
from database import engine, get_db, SessionLocal

# مسارات مطلقة تضمن قراءة قوالب الـ HTML ودعم التشغيل من أي مجلد
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
ADMIN_HTML_PATH = os.path.join(TEMPLATES_DIR, "admin.html")

# إنشاء الجداول في قاعدة البيانات تلقائياً عند أول تشغيل
models.Base.metadata.create_all(bind=engine)

app = FastAPI(
    title="Haris Anti-Fraud API & Admin Dashboard 🛡️",
    description="خادم ولوحة تحكم نظام حارس المتقدمة لإدارة الحظر ومراجعة البلاغات",
    version="1.0.0"
)

# السماح للاتصال من جميع المصادر (CORS)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

templates = Jinja2Templates(directory=TEMPLATES_DIR)

# بداية التشغيل بدون أي بيانات وهمية
@app.on_event("startup")
def startup_populate():
    pass

# ==================== 1. لوحة التحكم (Admin Dashboard UI) ====================

@app.get("/admin", response_class=HTMLResponse)
def get_admin_dashboard(request: Request):
    if os.path.exists(ADMIN_HTML_PATH):
        return FileResponse(ADMIN_HTML_PATH)
    return templates.TemplateResponse("admin.html", {"request": request})

@app.get("/", response_class=HTMLResponse)
def root_redirect(request: Request):
    if os.path.exists(ADMIN_HTML_PATH):
        return FileResponse(ADMIN_HTML_PATH)
    return templates.TemplateResponse("admin.html", {"request": request})

# ==================== 2. نظام مراجعة البلاغات (Moderation System) ====================

@app.post("/api/v1/reports/", response_model=schemas.ReportResponse)
def create_report(report: schemas.ReportCreate, db: Session = Depends(get_db)):
    """إرسال بلاغ جديد من التطبيق بحالة معلقة (status = pending)"""
    db_report = models.Report(
        phone_number=report.phone_number,
        threat_type=report.threat_type,
        report_reason=report.report_reason,
        status="pending"
    )
    db.add(db_report)
    db.commit()
    db.refresh(db_report)
    return db_report

@app.get("/api/v1/reports/", response_model=List[schemas.ReportResponse])
def list_reports(status: str = None, db: Session = Depends(get_db)):
    """عرض قائمة البلاغات مع إمكانية التصفية بحسب الحالة"""
    query = db.query(models.Report)
    if status:
        query = query.filter(models.Report.status == status)
    return query.order_by(models.Report.created_at.desc()).all()

@app.put("/api/v1/reports/{report_id}/approve", response_model=schemas.ReportResponse)
def approve_report(report_id: int, db: Session = Depends(get_db)):
    """الموافقة على البلاغ ونقل الرقم تلقائياً إلى قائمة الأرقام المحظورة النشطة"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    report.status = "approved"

    # نقل الرقم تلقائياً لقائمة الحظر المباشر إذا لم يكن موجوداً مسبقاً
    existing_blocked = db.query(models.BlockedNumber).filter(
        models.BlockedNumber.phone_number == report.phone_number
    ).first()

    if not existing_blocked:
        new_blocked = models.BlockedNumber(
            phone_number=report.phone_number,
            reason=f"تم حظره بناءً على موافقة البلاغ: {report.report_reason}",
            severity="خطر شديد"
        )
        db.add(new_blocked)

    db.commit()
    db.refresh(report)
    return report

@app.put("/api/v1/reports/{report_id}/reject", response_model=schemas.ReportResponse)
def reject_report(report_id: int, db: Session = Depends(get_db)):
    """رفض البلاغ الوارد"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    report.status = "rejected"
    db.commit()
    db.refresh(report)
    return report

@app.delete("/api/v1/reports/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    """حذف بلاغ محدد نهائياً من لوحة التحكم"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    db.delete(report)
    db.commit()
    return {"message": "تم حذف البلاغ بنجاح"}

# ==================== 3. إدارة الأرقام المشبوهة (Blocked Numbers CRUD & Import/Export) ====================

@app.get("/api/v1/blocked-numbers/export")
def export_blocked_numbers_excel(db: Session = Depends(get_db)):
    """تصدير كافة الأرقام المحظورة إلى ملف إكسيل (.xlsx) أو CSV"""
    numbers = db.query(models.BlockedNumber).order_by(models.BlockedNumber.created_at.desc()).all()

    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الأرقام المحظورة"

        # كتابة عناوين الأعمدة
        ws.append(["المعرف (ID)", "رقم الهاتف", "سبب الحظر", "درجة الخطورة", "تاريخ الحظر"])

        for num in numbers:
            date_str = num.created_at.strftime("%Y-%m-%d %H:%M:%S") if num.created_at else ""
            ws.append([num.id, num.phone_number, num.reason, num.severity, date_str])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        headers = {
            "Content-Disposition": "attachment; filename=haris_blocked_numbers.xlsx"
        }
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )
    except Exception as e:
        import csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Phone Number", "Reason", "Severity", "Date Created"])
        for num in numbers:
            date_str = num.created_at.strftime("%Y-%m-%d %H:%M:%S") if num.created_at else ""
            writer.writerow([num.id, num.phone_number, num.reason, num.severity, date_str])

        headers = {
            "Content-Disposition": "attachment; filename=haris_blocked_numbers.csv"
        }
        return Response(
            content=output.getvalue().encode("utf-8-sig"),
            media_type="text/csv",
            headers=headers
        )

@app.post("/api/v1/blocked-numbers/import")
async def import_blocked_numbers_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """رفع واستيراد أرقام محظورة من ملف إكسيل (.xlsx / .csv)"""
    filename = file.filename.lower()
    contents = await file.read()

    imported_count = 0
    skipped_count = 0
    rows = []

    if filename.endswith(".csv"):
        import csv
        text = contents.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        for r in reader:
            if r: rows.append(r)
    else:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents))
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if row:
                    rows.append([str(c) if c is not None else "" for c in row])
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"تعذر قراءة ملف الإكسيل: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="الملف المرفوع فارغ")

    header_skipped = False
    for row in rows:
        if len(row) == 0: continue
        phone = str(row[0]).strip()

        if not header_skipped and ("رقم" in phone.lower() or "phone" in phone.lower() or "id" in phone.lower()):
            header_skipped = True
            continue

        if not phone: continue

        reason = str(row[1]).strip() if len(row) > 1 and row[1] else "تم الاستيراد من ملف خارجي"
        severity = str(row[2]).strip() if len(row) > 2 and row[2] else "خطر شديد"

        # فحص وجود الرقم لمنع التكرار
        existing = db.query(models.BlockedNumber).filter(
            models.BlockedNumber.phone_number == phone
        ).first()

        if not existing:
            db_num = models.BlockedNumber(
                phone_number=phone,
                reason=reason,
                severity=severity
            )
            db.add(db_num)
            imported_count += 1
        else:
            skipped_count += 1

    db.commit()
    return {
        "imported_count": imported_count,
        "skipped_count": skipped_count,
        "message": f"تم استيراد {imported_count} رقم بنجاح! وتجاوز {skipped_count} رقم مكرر."
    }

@app.get("/api/v1/blocked-numbers/", response_model=List[schemas.BlockedNumberResponse])
def get_blocked_numbers(db: Session = Depends(get_db)):
    """جلب كافة الأرقام المحظورة النشطة"""
    return db.query(models.BlockedNumber).order_by(models.BlockedNumber.created_at.desc()).all()

@app.post("/api/v1/blocked-numbers/", response_model=schemas.BlockedNumberResponse)
def add_blocked_number(number: schemas.BlockedNumberCreate, db: Session = Depends(get_db)):
    """إضافة رقم محظور جديد يدوياً من لوحة التحكم"""
    existing = db.query(models.BlockedNumber).filter(
        models.BlockedNumber.phone_number == number.phone_number
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="الرقم محظور بالفعل سابقاً")

    db_number = models.BlockedNumber(
        phone_number=number.phone_number,
        reason=number.reason,
        severity=number.severity
    )
    db.add(db_number)
    db.commit()
    db.refresh(db_number)
    return db_number

@app.put("/api/v1/blocked-numbers/{number_id}", response_model=schemas.BlockedNumberResponse)
def update_blocked_number(number_id: int, update_data: schemas.BlockedNumberUpdate, db: Session = Depends(get_db)):
    """تعديل بيانات رقم محظور في لوحة التحكم"""
    db_number = db.query(models.BlockedNumber).filter(models.BlockedNumber.id == number_id).first()
    if not db_number:
        raise HTTPException(status_code=404, detail="الرقم غير موجود")

    if update_data.phone_number is not None:
        db_number.phone_number = update_data.phone_number
    if update_data.reason is not None:
        db_number.reason = update_data.reason
    if update_data.severity is not None:
        db_number.severity = update_data.severity

    db.commit()
    db.refresh(db_number)
    return db_number

@app.delete("/api/v1/blocked-numbers/{number_id}")
def delete_blocked_number(number_id: int, db: Session = Depends(get_db)):
    """حذف رقم من قائمة الحظر من لوحة التحكم"""
    db_number = db.query(models.BlockedNumber).filter(models.BlockedNumber.id == number_id).first()
    if not db_number:
        raise HTTPException(status_code=404, detail="الرقم غير موجود")

    db.delete(db_number)
    db.commit()
    return {"message": "تم حذف الرقم من قائمة الحظر بنجاح"}

# ==================== 4. استعلامات تطبيق الأندرويد السريعة ====================

@app.get("/api/v1/check-number")
def check_number(phone: str, db: Session = Depends(get_db)):
    """استعلام ذكي سريع لتطبيق الأندرويد للتحقق مما إذا كان الرقم محظوراً أم لا بنوعين مقارنة طردية ونظيفة"""
    clean_input = phone.replace(" ", "").replace("-", "").replace("+", "")

    all_blocked = db.query(models.BlockedNumber).all()
    for b in all_blocked:
        b_clean = b.phone_number.replace(" ", "").replace("-", "").replace("+", "")
        if clean_input in b_clean or b_clean in clean_input or b.phone_number == phone:
            return {
                "is_blocked": True,
                "phone_number": b.phone_number,
                "reason": b.reason,
                "severity": b.severity
            }

    return {"is_blocked": False, "phone_number": phone}
