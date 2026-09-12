import os
import io
from typing import List
from fastapi import FastAPI, Depends, HTTPException, Request, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, Response
from sqlalchemy.orm import Session

import models
import schemas
from database import engine, get_db, SessionLocal

# مسارات مطلقة تضمن قراءة قوالب الـ HTML ودعم التشغيل من أي مجلد برمجياً
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
ADMIN_HTML_PATH = os.path.join(TEMPLATES_DIR, "admin.html")

# إنشاء الجداول في قاعدة البيانات تلقائياً عند أول تشغيل
models.Base.metadata.create_all(bind=engine)

app = FastAPI(
    title="Haris Anti-Fraud API & Admin Dashboard 🛡️",
    description="خادم ولوحة تحكم نظام حارس المتقدمة لإدارة الحظر ومراجعة البلاغات",
    version="1.0.0"
)

# السماح للاتصال من جميع المصادر (CORS)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# بداية التشغيل بدون أي بيانات وهمية
@app.on_event("startup")
def startup_populate():
    pass

# ==================== 1. لوحة التحكم (Admin Dashboard UI) ====================

DEFAULT_ADMIN_HTML = """<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>لوحة تحكم تطبيق حارس 🛡️</title>
    <!-- FontAwesome & Google Fonts -->
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
    <link href="https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;800&display=swap" rel="stylesheet">
    <style>
        :root {
            --bg-color: #0A0E1A;
            --card-bg: #131B2E;
            --accent-cyan: #00E5FF;
            --accent-green: #10B981;
            --accent-red: #EF4444;
            --accent-orange: #F59E0B;
            --border-color: #1E293B;
            --text-main: #FFFFFF;
            --text-sub: #94A3B8;
        }

        * {
            box-sizing: border-box;
            margin: 0;
            padding: 0;
            font-family: 'Cairo', sans-serif;
        }

        body {
            background-color: var(--bg-color);
            color: var(--text-main);
            padding: 24px;
            direction: rtl;
        }

        /* Top Bar */
        .top-bar {
            display: flex;
            justify-content: space-between;
            align-items: center;
            background-color: var(--card-bg);
            padding: 16px 24px;
            border-radius: 16px;
            border: 1px solid var(--border-color);
            margin-bottom: 24px;
        }

        .brand {
            display: flex;
            align-items: center;
            gap: 12px;
        }

        .brand i {
            font-size: 28px;
            color: var(--accent-cyan);
        }

        .brand h1 {
            font-size: 20px;
            font-weight: 800;
        }

        /* Stats Grid */
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
            gap: 16px;
            margin-bottom: 24px;
        }

        .stat-card {
            background-color: var(--card-bg);
            padding: 20px;
            border-radius: 16px;
            border: 1px solid var(--border-color);
            display: flex;
            flex-direction: column;
            align-items: center;
            text-align: center;
        }

        .stat-value {
            font-size: 32px;
            font-weight: 800;
            margin-bottom: 4px;
        }

        .stat-label {
            font-size: 14px;
            color: var(--text-sub);
        }

        /* Section Containers */
        .section-container {
            background-color: var(--card-bg);
            border-radius: 16px;
            border: 1px solid var(--border-color);
            padding: 24px;
            margin-bottom: 24px;
        }

        .section-header {
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 20px;
            flex-wrap: wrap;
            gap: 12px;
        }

        .section-title {
            font-size: 18px;
            font-weight: 700;
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .section-title i {
            color: var(--accent-cyan);
        }

        /* Buttons */
        .btn {
            padding: 8px 16px;
            border-radius: 10px;
            font-weight: 700;
            font-size: 14px;
            border: none;
            cursor: pointer;
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }

        .btn-cyan {
            background-color: var(--accent-cyan);
            color: #0A0E1A;
        }

        .btn-cyan:hover {
            opacity: 0.9;
            transform: translateY(-1px);
        }

        .btn-green {
            background-color: var(--accent-green);
            color: #FFFFFF;
        }

        .btn-green:hover {
            opacity: 0.9;
            transform: translateY(-1px);
        }

        .btn-approve {
            background-color: rgba(16, 185, 129, 0.15);
            color: var(--accent-green);
            border: 1px solid var(--accent-green);
        }

        .btn-approve:hover {
            background-color: var(--accent-green);
            color: #FFFFFF;
        }

        .btn-reject {
            background-color: rgba(239, 68, 68, 0.15);
            color: var(--accent-red);
            border: 1px solid var(--accent-red);
        }

        .btn-reject:hover {
            background-color: var(--accent-red);
            color: #FFFFFF;
        }

        .btn-delete {
            background-color: rgba(239, 68, 68, 0.15);
            color: var(--accent-red);
            border: 1px solid var(--accent-red);
        }

        .btn-delete:hover {
            background-color: var(--accent-red);
            color: #FFFFFF;
        }

        .btn-edit {
            background-color: rgba(245, 158, 11, 0.15);
            color: var(--accent-orange);
            border: 1px solid var(--accent-orange);
        }

        .btn-edit:hover {
            background-color: var(--accent-orange);
            color: #FFFFFF;
        }

        /* Tables */
        .table-responsive {
            overflow-x: auto;
        }

        table {
            width: 100%;
            border-collapse: collapse;
            text-align: right;
        }

        th, td {
            padding: 14px 16px;
            border-bottom: 1px solid var(--border-color);
            font-size: 14px;
        }

        th {
            color: var(--text-sub);
            font-weight: 600;
            background-color: rgba(10, 14, 26, 0.5);
        }

        tbody tr:hover {
            background-color: rgba(255, 255, 255, 0.02);
        }

        /* Status Badges */
        .badge {
            padding: 4px 10px;
            border-radius: 50px;
            font-size: 12px;
            font-weight: 700;
            display: inline-block;
        }

        .badge-pending {
            background-color: rgba(245, 158, 11, 0.15);
            color: var(--accent-orange);
            border: 1px solid var(--accent-orange);
        }

        .badge-approved {
            background-color: rgba(16, 185, 129, 0.15);
            color: var(--accent-green);
            border: 1px solid var(--accent-green);
        }

        .badge-rejected {
            background-color: rgba(239, 68, 68, 0.15);
            color: var(--accent-red);
            border: 1px solid var(--accent-red);
        }

        /* Modal */
        .modal-overlay {
            display: none;
            position: fixed;
            top: 0;
            left: 0;
            width: 100%;
            height: 100%;
            background-color: rgba(0, 0, 0, 0.75);
            z-index: 1000;
            justify-content: center;
            align-items: center;
        }

        .modal-card {
            background-color: var(--card-bg);
            border-radius: 16px;
            border: 1px solid var(--border-color);
            width: 90%;
            max-width: 480px;
            padding: 24px;
        }

        .modal-title {
            font-size: 18px;
            font-weight: 700;
            margin-bottom: 16px;
        }

        .form-group {
            margin-bottom: 16px;
        }

        .form-group label {
            display: block;
            margin-bottom: 6px;
            font-size: 14px;
            color: var(--text-sub);
        }

        .form-control {
            width: 100%;
            padding: 10px 14px;
            background-color: var(--bg-color);
            border: 1px solid var(--border-color);
            border-radius: 10px;
            color: var(--text-main);
            font-size: 14px;
            outline: none;
        }

        .form-control:focus {
            border-color: var(--accent-cyan);
        }

        .modal-actions {
            display: flex;
            justify-content: flex-end;
            gap: 10px;
            margin-top: 20px;
        }
    </style>
</head>
<body>

    <!-- Top Bar -->
    <div class="top-bar">
        <div class="brand">
            <i class="fa-solid fa-shield-halved"></i>
            <h1>لوحة تحكم إدارة نظام حارس 🛡️</h1>
        </div>
        <div>
            <span style="color: var(--text-sub); font-size: 14px;">خادم FastAPI نشط ✅</span>
        </div>
    </div>

    <!-- Analytics Cards -->
    <div class="stats-grid">
        <div class="stat-card">
            <div class="stat-value" id="statBlockedCount" style="color: var(--accent-cyan);">0</div>
            <div class="stat-label">الأرقام المحظورة النشطة</div>
        </div>
        <div class="stat-card">
            <div class="stat-value" id="statPendingCount" style="color: var(--accent-orange);">0</div>
            <div class="stat-label">البلاغات المعلقة (Pending)</div>
        </div>
        <div class="stat-card">
            <div class="stat-value" id="statApprovedCount" style="color: var(--accent-green);">0</div>
            <div class="stat-label">البلاغات المقبولة</div>
        </div>
        <div class="stat-card">
            <div class="stat-value" id="statRejectedCount" style="color: var(--accent-red);">0</div>
            <div class="stat-label">البلاغات المرفوضة</div>
        </div>
    </div>

    <!-- Section 1: Moderation System -->
    <div class="section-container">
        <div class="section-header">
            <div class="section-title">
                <i class="fa-solid fa-user-check"></i>
                <span>نظام مراجعة البلاغات الواردة (Moderation System)</span>
            </div>
            <button class="btn btn-cyan" onclick="loadReports()"><i class="fa-solid fa-rotate-right"></i> تحديث القائمة</button>
        </div>

        <div class="table-responsive">
            <table>
                <thead>
                    <tr>
                        <th>المعرف</th>
                        <th>رقم الهاتف المشبوه</th>
                        <th>تكرار البلاغات</th>
                        <th>نوع التهديد</th>
                        <th>تفاصيل/سبب البلاغ</th>
                        <th>الحالة</th>
                        <th>تاريخ البلاغ</th>
                        <th>إجراء المراجعة</th>
                    </tr>
                </thead>
                <tbody id="reportsTableBody">
                    <tr><td colspan="8" style="text-align:center;">جاري تحميل البلاغات...</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <!-- Section 2: Blocked Numbers CRUD & Import/Export -->
    <div class="section-container">
        <div class="section-header">
            <div class="section-title">
                <i class="fa-solid fa-ban"></i>
                <span>إدارة الأرقام المشبوهة المحظورة (Active Blocked Numbers)</span>
            </div>
            <div style="display: flex; gap: 8px; flex-wrap: wrap;">
                <button class="btn btn-green" onclick="downloadExcel()"><i class="fa-solid fa-file-excel"></i> تنزيل القائمة (تصدير إكسيل)</button>
                <button class="btn btn-cyan" onclick="openImportModal()"><i class="fa-solid fa-file-import"></i> رفع ملف إكسيل/CSV</button>
                <button class="btn btn-cyan" onclick="openAddModal()"><i class="fa-solid fa-plus"></i> إضافة رقم يدوياً</button>
            </div>
        </div>

        <div class="table-responsive">
            <table>
                <thead>
                    <tr>
                        <th>المعرف</th>
                        <th>رقم الهاتف</th>
                        <th>سبب الحظر</th>
                        <th>درجة الخطورة</th>
                        <th>تاريخ الحظر</th>
                        <th>الإجراءات</th>
                    </tr>
                </thead>
                <tbody id="blockedTableBody">
                    <tr><td colspan="6" style="text-align:center;">جاري تحميل الأرقام المحظورة...</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <!-- Add/Edit Number Modal -->
    <div class="modal-overlay" id="numberModal">
        <div class="modal-card">
            <div class="modal-title" id="modalTitle">إضافة رقم محظور جديد</div>
            <form id="numberForm" onsubmit="handleFormSubmit(event)">
                <input type="hidden" id="editNumberId">
                <div class="form-group">
                    <label>رقم الهاتف المشبوه</label>
                    <input type="text" id="inputPhoneNumber" class="form-control" placeholder="+966 50 XXX XXXX" required>
                </div>
                <div class="form-group">
                    <label>سبب الحظر</label>
                    <input type="text" id="inputReason" class="form-control" placeholder="مثال: رسائل احتيال وتصيد بنكي" required>
                </div>
                <div class="form-group">
                    <label>درجة الخطورة</label>
                    <select id="inputSeverity" class="form-control">
                        <option value="خطر شديد">خطر شديد ⚠️</option>
                        <option value="رقم مشبوه">رقم مشبوه 🟠</option>
                        <option value="خطر احتيال">خطر احتيال 🚨</option>
                    </select>
                </div>
                <div class="modal-actions">
                    <button type="button" class="btn btn-reject" onclick="closeModal()">إلغاء</button>
                    <button type="submit" class="btn btn-cyan">حفظ البيانات</button>
                </div>
            </form>
        </div>
    </div>

    <!-- Import Excel Modal -->
    <div class="modal-overlay" id="importModal">
        <div class="modal-card">
            <div class="modal-title">رفع واستيراد ملف إكسيل / CSV بالأرقام المحظورة</div>
            <form id="importForm" onsubmit="handleImportSubmit(event)">
                <div class="form-group">
                    <label>اختر ملف الإكسيل (.xlsx / .csv)</label>
                    <input type="file" id="importFile" class="form-control" accept=".xlsx, .xls, .csv" required>
                    <p style="font-size: 12px; color: var(--text-sub); margin-top: 8px;">* يجب أن يضم الملف الأعمدة: (رقم الهاتف، سبب الحظر، درجة الخطورة).</p>
                </div>
                <div class="modal-actions">
                    <button type="button" class="btn btn-reject" onclick="closeImportModal()">إلغاء</button>
                    <button type="submit" class="btn btn-green"><i class="fa-solid fa-upload"></i> رفع واستيراد البيانات</button>
                </div>
            </form>
        </div>
    </div>

    <!-- Dashboard Logic JS -->
    <script>
        const API_BASE = '/api/v1';

        document.addEventListener('DOMContentLoaded', () => {
            loadDashboardData();
        });

        async function loadDashboardData() {
            await Promise.all([loadReports(), loadBlockedNumbers()]);
        }

        // Load Reports (Moderation)
        async function loadReports() {
            try {
                const res = await fetch(`${API_BASE}/reports/`);
                const reports = await res.json();

                const tbody = document.getElementById('reportsTableBody');
                tbody.innerHTML = '';

                let pending = 0, approved = 0, rejected = 0;

                if (reports.length === 0) {
                    tbody.innerHTML = '<tr><td colspan="8" style="text-align:center;">لا توجد بلاغات مسجلة حالياً.</td></tr>';
                } else {
                    reports.forEach(report => {
                        if (report.status === 'pending') pending++;
                        else if (report.status === 'approved') approved++;
                        else if (report.status === 'rejected') rejected++;

                        let badgeClass = 'badge-pending';
                        let badgeText = 'معلق (Pending)';
                        if (report.status === 'approved') {
                            badgeClass = 'badge-approved';
                            badgeText = 'مقبول (Approved)';
                        } else if (report.status === 'rejected') {
                            badgeClass = 'badge-rejected';
                            badgeText = 'مرفوض (Rejected)';
                        }

                        let countBadge = `<span class="badge badge-pending">${report.reports_count || 1} بلاغ</span>`;
                        if (report.reports_count > 1) {
                            countBadge = `<span class="badge badge-rejected" style="background: rgba(239, 68, 68, 0.2); color: #EF4444;"><i class="fa-solid fa-fire"></i> ${report.reports_count} بلاغات متكررة</span>`;
                        }

                        const dateStr = new Date(report.created_at).toLocaleString('ar-SA');

                        const tr = document.createElement('tr');
                        tr.innerHTML = `
                            <td>${report.id}</td>
                            <td><strong>${report.phone_number}</strong></td>
                            <td>${countBadge}</td>
                            <td>${report.threat_type}</td>
                            <td>${report.report_reason}</td>
                            <td><span class="badge ${badgeClass}">${badgeText}</span></td>
                            <td>${dateStr}</td>
                            <td>
                                ${report.status === 'pending' ? `
                                    <button class="btn btn-approve" onclick="approveReport(${report.id})"><i class="fa-solid fa-check"></i> موافقة</button>
                                    <button class="btn btn-reject" onclick="rejectReport(${report.id})"><i class="fa-solid fa-xmark"></i> رفض</button>
                                ` : '<span style="color:var(--text-sub); margin-left:6px;">تمت المراجعة</span>'}
                                <button class="btn btn-delete" onclick="deleteReport(${report.id})"><i class="fa-solid fa-trash"></i> حذف</button>
                            </td>
                        `;
                        tbody.appendChild(tr);
                    });
                }

                document.getElementById('statPendingCount').innerText = pending;
                document.getElementById('statApprovedCount').innerText = approved;
                document.getElementById('statRejectedCount').innerText = rejected;

            } catch (err) {
                console.error(err);
            }
        }

        // Approve Report
        async function approveReport(id) {
            if (!confirm('هل أنت ألكيد من الموافقة على البلاغ ونقل الرقم تلقائياً لقائمة الحظر النشطة؟')) return;
            try {
                const res = await fetch(`${API_BASE}/reports/${id}/approve`, { method: 'PUT' });
                if (res.ok) {
                    alert('تمت الموافقة ونقل الرقم لقائمة الحظر بنجاح 🛡️');
                    loadDashboardData();
                }
            } catch (err) {
                alert('حدث خطأ أثناء الموافقة');
            }
        }

        // Reject Report
        async function rejectReport(id) {
            if (!confirm('هل أنت ألكيد من رفض البلاغ؟')) return;
            try {
                const res = await fetch(`${API_BASE}/reports/${id}/reject`, { method: 'PUT' });
                if (res.ok) {
                    alert('تم رفض البلاغ بنجاح');
                    loadDashboardData();
                }
            } catch (err) {
                alert('حدث خطأ أثناء رفض البلاغ');
            }
        }

        // Delete Report
        async function deleteReport(id) {
            if (!confirm('هل أنت ألكيد من حذف هذا البلاغ نهائياً؟')) return;
            try {
                const res = await fetch(`${API_BASE}/reports/${id}`, { method: 'DELETE' });
                if (res.ok) {
                    loadDashboardData();
                }
            } catch (err) {
                alert('حدث خطأ أثناء حذف البلاغ');
            }
        }

        // Download Excel
        function downloadExcel() {
            window.location.href = `${API_BASE}/blocked-numbers/export`;
        }

        // Open/Close Import Modal
        function openImportModal() {
            document.getElementById('importFile').value = '';
            document.getElementById('importModal').style.display = 'flex';
        }

        function closeImportModal() {
            document.getElementById('importModal').style.display = 'none';
        }

        async function handleImportSubmit(e) {
            e.preventDefault();
            const fileInput = document.getElementById('importFile');
            if (!fileInput.files || fileInput.files.length === 0) {
                alert('يرجى اختيار ملف الإكسيل أولاً');
                return;
            }

            const formData = new FormData();
            formData.append('file', fileInput.files[0]);

            try {
                const res = await fetch(`${API_BASE}/blocked-numbers/import`, {
                    method: 'POST',
                    body: formData
                });
                const result = await res.json();
                if (res.ok) {
                    alert(result.message);
                    closeImportModal();
                    loadBlockedNumbers();
                } else {
                    alert('خطأ في رفع الملف: ' + (result.detail || 'يرجى التأكد من صيغة الملف'));
                }
            } catch (err) {
                alert('حدث خطأ أثناء رفع الملف للسيرفر');
            }
        }

        // Load Blocked Numbers
        async function loadBlockedNumbers() {
            try {
                const res = await fetch(`${API_BASE}/blocked-numbers/`);
                const numbers = await res.json();

                const tbody = document.getElementById('blockedTableBody');
                tbody.innerHTML = '';

                document.getElementById('statBlockedCount').innerText = numbers.length;

                if (numbers.length === 0) {
                    tbody.innerHTML = '<tr><td colspan="6" style="text-align:center;">لا توجد أرقام محظورة نشطة.</td></tr>';
                } else {
                    numbers.forEach(num => {
                        const dateStr = new Date(num.created_at).toLocaleString('ar-SA');
                        const tr = document.createElement('tr');
                        tr.innerHTML = `
                            <td>${num.id}</td>
                            <td><strong>${num.phone_number}</strong></td>
                            <td>${num.reason}</td>
                            <td><span class="badge badge-rejected">${num.severity}</span></td>
                            <td>${dateStr}</td>
                            <td>
                                <button class="btn btn-edit" onclick="openEditModal(${num.id}, '${num.phone_number}', '${num.reason}', '${num.severity}')"><i class="fa-solid fa-pen"></i> تعديل</button>
                                <button class="btn btn-delete" onclick="deleteBlockedNumber(${num.id})"><i class="fa-solid fa-trash"></i> حذف</button>
                            </td>
                        `;
                        tbody.appendChild(tr);
                    });
                }
            } catch (err) {
                console.error(err);
            }
        }

        // Delete Blocked Number
        async function deleteBlockedNumber(id) {
            if (!confirm('هل أنت ألكيد من حذف هذا الرقم من قائمة الحظر؟')) return;
            try {
                const res = await fetch(`${API_BASE}/blocked-numbers/${id}`, { method: 'DELETE' });
                if (res.ok) {
                    loadBlockedNumbers();
                }
            } catch (err) {
                alert('حدث خطأ أثناء الحذف');
            }
        }

        // Modal Open / Close
        function openAddModal() {
            document.getElementById('editNumberId').value = '';
            document.getElementById('inputPhoneNumber').value = '';
            document.getElementById('inputReason').value = '';
            document.getElementById('inputSeverity').value = 'خطر شديد';
            document.getElementById('modalTitle').innerText = 'إضافة رقم محظور جديد';
            document.getElementById('numberModal').style.display = 'flex';
        }

        function openEditModal(id, phone, reason, severity) {
            document.getElementById('editNumberId').value = id;
            document.getElementById('inputPhoneNumber').value = phone;
            document.getElementById('inputReason').value = reason;
            document.getElementById('inputSeverity').value = severity;
            document.getElementById('modalTitle').innerText = 'تعديل بيانات الرقم المحظور';
            document.getElementById('numberModal').style.display = 'flex';
        }

        function closeModal() {
            document.getElementById('numberModal').style.display = 'none';
        }

        async function handleFormSubmit(e) {
            e.preventDefault();
            const id = document.getElementById('editNumberId').value;
            const phone_number = document.getElementById('inputPhoneNumber').value;
            const reason = document.getElementById('inputReason').value;
            const severity = document.getElementById('inputSeverity').value;

            const payload = { phone_number, reason, severity };

            try {
                let res;
                if (id) {
                    res = await fetch(`${API_BASE}/blocked-numbers/${id}`, {
                        method: 'PUT',
                        headers: { 'Content-Type': 'application/json' },
                        body: JSON.stringify(payload)
                    });
                } else {
                    res = await fetch(`${API_BASE}/blocked-numbers/`, {
                        method: 'POST',
                        headers: { 'Content-Type': 'application/json' },
                        body: JSON.stringify(payload)
                    });
                }

                if (res.ok) {
                    closeModal();
                    loadBlockedNumbers();
                } else {
                    alert('تعذر الحفظ، يرجى التأكد من أن الرقم غير مكرر');
                }
            } catch (err) {
                alert('حدث خطأ في الاتصال بالخادم');
            }
        }
    </script>
</body>
</html>"""

def render_admin_html() -> str:
    """قراءة وعرض صفحة لوحة التحكم بأمان كامل بدون أي أخطاء سيرفر 500"""
    try:
        if os.path.exists(ADMIN_HTML_PATH):
            with open(ADMIN_HTML_PATH, "r", encoding="utf-8") as f:
                return f.read()
    except Exception as e:
        print(f"Error reading admin.html: {e}")

    return DEFAULT_ADMIN_HTML

@app.get("/admin", response_class=HTMLResponse)
def get_admin_dashboard():
    return HTMLResponse(content=render_admin_html())

@app.get("/", response_class=HTMLResponse)
def root_redirect():
    return HTMLResponse(content=render_admin_html())

# ==================== 2. نظام مراجعة البلاغات (Moderation System) ====================

@app.post("/api/v1/reports/", response_model=schemas.ReportResponse)
def create_report(report: schemas.ReportCreate, db: Session = Depends(get_db)):
    """إرسال بلاغ جديد من التطبيق بحالة معلقة (status = pending)"""
    db_report = models.Report(
        phone_number=report.phone_number,
        threat_type=report.threat_type,
        report_reason=report.report_reason,
        status="pending"
    )
    db.add(db_report)
    db.commit()
    db.refresh(db_report)
    return db_report

@app.get("/api/v1/reports/", response_model=List[schemas.ReportResponse])
def list_reports(status: str = None, db: Session = Depends(get_db)):
    """عرض قائمة البلاغات مع احتساب وتأكيد تكرار البلاغات المقدمة على كل رقم هاتف"""
    query = db.query(models.Report)
    if status:
        query = query.filter(models.Report.status == status)

    reports = query.order_by(models.Report.created_at.desc()).all()

    result = []
    for r in reports:
        # حساب إجمالي عدد البلاغات المقدمة على هذا الرقم بعينه
        count = db.query(models.Report).filter(
            models.Report.phone_number == r.phone_number
        ).count()

        rep_dict = schemas.ReportResponse(
            id=r.id,
            phone_number=r.phone_number,
            threat_type=r.threat_type,
            report_reason=r.report_reason,
            status=r.status,
            reports_count=count,
            created_at=r.created_at
        )
        result.append(rep_dict)

    return result

@app.put("/api/v1/reports/{report_id}/approve", response_model=schemas.ReportResponse)
def approve_report(report_id: int, db: Session = Depends(get_db)):
    """الموافقة على البلاغ ونقل الرقم تلقائياً إلى قائمة الأرقام المحظورة النشطة"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    report.status = "approved"

    # نقل الرقم تلقائياً لقائمة الحظر المباشر إذا لم يكن موجوداً مسبقاً
    existing_blocked = db.query(models.BlockedNumber).filter(
        models.BlockedNumber.phone_number == report.phone_number
    ).first()

    if not existing_blocked:
        new_blocked = models.BlockedNumber(
            phone_number=report.phone_number,
            reason=f"تم حظره بناءً على موافقة البلاغ: {report.report_reason}",
            severity="خطر شديد"
        )
        db.add(new_blocked)

    db.commit()
    db.refresh(report)
    return report

@app.put("/api/v1/reports/{report_id}/reject", response_model=schemas.ReportResponse)
def reject_report(report_id: int, db: Session = Depends(get_db)):
    """رفض البلاغ الوارد"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    report.status = "rejected"
    db.commit()
    db.refresh(report)
    return report

@app.delete("/api/v1/reports/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    """حذف بلاغ محدد نهائياً من لوحة التحكم"""
    report = db.query(models.Report).filter(models.Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="البلاغ غير موجود")

    db.delete(report)
    db.commit()
    return {"message": "تم حذف البلاغ بنجاح"}

# ==================== 3. إدارة الأرقام المشبوهة (Blocked Numbers CRUD & Import/Export) ====================

@app.get("/api/v1/blocked-numbers/export")
def export_blocked_numbers_excel(db: Session = Depends(get_db)):
    """تصدير كافة الأرقام المحظورة إلى ملف إكسيل (.xlsx) أو CSV"""
    numbers = db.query(models.BlockedNumber).order_by(models.BlockedNumber.created_at.desc()).all()

    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الأرقام المحظورة"

        # كتابة عناوين الأعمدة
        ws.append(["المعرف (ID)", "رقم الهاتف", "سبب الحظر", "درجة الخطورة", "تاريخ الحظر"])

        for num in numbers:
            date_str = num.created_at.strftime("%Y-%m-%d %H:%M:%S") if num.created_at else ""
            ws.append([num.id, num.phone_number, num.reason, num.severity, date_str])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        headers = {
            "Content-Disposition": "attachment; filename=haris_blocked_numbers.xlsx"
        }
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers=headers
        )
    except Exception as e:
        import csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Phone Number", "Reason", "Severity", "Date Created"])
        for num in numbers:
            date_str = num.created_at.strftime("%Y-%m-%d %H:%M:%S") if num.created_at else ""
            writer.writerow([num.id, num.phone_number, num.reason, num.severity, date_str])

        headers = {
            "Content-Disposition": "attachment; filename=haris_blocked_numbers.csv"
        }
        return Response(
            content=output.getvalue().encode("utf-8-sig"),
            media_type="text/csv",
            headers=headers
        )

@app.post("/api/v1/blocked-numbers/import")
async def import_blocked_numbers_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """رفع واستيراد أرقام محظورة من ملف إكسيل (.xlsx / .csv)"""
    filename = file.filename.lower()
    contents = await file.read()

    imported_count = 0
    skipped_count = 0
    rows = []

    if filename.endswith(".csv"):
        import csv
        text = contents.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        for r in reader:
            if r: rows.append(r)
    else:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents))
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if row:
                    rows.append([str(c) if c is not None else "" for c in row])
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"تعذر قراءة ملف الإكسيل: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="الملف المرفوع فارغ")

    header_skipped = False
    for row in rows:
        if len(row) == 0: continue
        phone = str(row[0]).strip()

        if not header_skipped and ("رقم" in phone.lower() or "phone" in phone.lower() or "id" in phone.lower()):
            header_skipped = True
            continue

        if not phone: continue

        reason = str(row[1]).strip() if len(row) > 1 and row[1] else "تم الاستيراد من ملف خارجي"
        severity = str(row[2]).strip() if len(row) > 2 and row[2] else "خطر شديد"

        # فحص وجود الرقم لمنع التكرار
        existing = db.query(models.BlockedNumber).filter(
            models.BlockedNumber.phone_number == phone
        ).first()

        if not existing:
            db_num = models.BlockedNumber(
                phone_number=phone,
                reason=reason,
                severity=severity
            )
            db.add(db_num)
            imported_count += 1
        else:
            skipped_count += 1

    db.commit()
    return {
        "imported_count": imported_count,
        "skipped_count": skipped_count,
        "message": f"تم استيراد {imported_count} رقم بنجاح! وتجاوز {skipped_count} رقم مكرر."
    }

@app.get("/api/v1/blocked-numbers/", response_model=List[schemas.BlockedNumberResponse])
def get_blocked_numbers(db: Session = Depends(get_db)):
    """جلب كافة الأرقام المحظورة النشطة"""
    return db.query(models.BlockedNumber).order_by(models.BlockedNumber.created_at.desc()).all()

@app.post("/api/v1/blocked-numbers/", response_model=schemas.BlockedNumberResponse)
def add_blocked_number(number: schemas.BlockedNumberCreate, db: Session = Depends(get_db)):
    """إضافة رقم محظور جديد يدوياً من لوحة التحكم"""
    existing = db.query(models.BlockedNumber).filter(
        models.BlockedNumber.phone_number == number.phone_number
    ).first()
    if existing:
        raise HTTPException(status_code=400, detail="الرقم محظور بالفعل سابقاً")

    db_number = models.BlockedNumber(
        phone_number=number.phone_number,
        reason=number.reason,
        severity=number.severity
    )
    db.add(db_number)
    db.commit()
    db.refresh(db_number)
    return db_number

@app.put("/api/v1/blocked-numbers/{number_id}", response_model=schemas.BlockedNumberResponse)
def update_blocked_number(number_id: int, update_data: schemas.BlockedNumberUpdate, db: Session = Depends(get_db)):
    """تعديل بيانات رقم محظور في لوحة التحكم"""
    db_number = db.query(models.BlockedNumber).filter(models.BlockedNumber.id == number_id).first()
    if not db_number:
        raise HTTPException(status_code=404, detail="الرقم غير موجود")

    if update_data.phone_number is not None:
        db_number.phone_number = update_data.phone_number
    if update_data.reason is not None:
        db_number.reason = update_data.reason
    if update_data.severity is not None:
        db_number.severity = update_data.severity

    db.commit()
    db.refresh(db_number)
    return db_number

@app.delete("/api/v1/blocked-numbers/{number_id}")
def delete_blocked_number(number_id: int, db: Session = Depends(get_db)):
    """حذف رقم من قائمة الحظر من لوحة التحكم"""
    db_number = db.query(models.BlockedNumber).filter(models.BlockedNumber.id == number_id).first()
    if not db_number:
        raise HTTPException(status_code=404, detail="الرقم غير موجود")

    db.delete(db_number)
    db.commit()
    return {"message": "تم حذف الرقم من قائمة الحظر بنجاح"}

# ==================== 4. استعلامات تطبيق الأندرويد السريعة ====================

@app.get("/api/v1/check-number")
def check_number(phone: str, db: Session = Depends(get_db)):
    """استعلام ذكي سريع لتطبيق الأندرويد للتحقق مما إذا كان الرقم محظوراً أم لا بنوعين مقارنة طردية ونظيفة"""
    clean_input = phone.replace(" ", "").replace("-", "").replace("+", "")

    all_blocked = db.query(models.BlockedNumber).all()
    for b in all_blocked:
        b_clean = b.phone_number.replace(" ", "").replace("-", "").replace("+", "")
        if clean_input in b_clean or b_clean in clean_input or b.phone_number == phone:
            return {
                "is_blocked": True,
                "phone_number": b.phone_number,
                "reason": b.reason,
                "severity": b.severity
            }

    return {"is_blocked": False, "phone_number": phone}

# ==================== 5. تحليل الذكاء الاصطناعي عبر Ollama / Llama 3 ====================

class MessageAnalysisRequest(schemas.BaseModel):
    message_body: str
    sender: str = ""

@app.post("/api/v1/analyze-message")
async def analyze_message_ai(req: MessageAnalysisRequest):
    """تحليل فوري لسياق الرسالة باستخدام موديل Llama 3 للذكاء الاصطناعي"""
    prompt = f"""You are Haris AI Anti-Phishing Fraud Security Classifier.
Analyze the following text (SMS, WhatsApp message, or Call transcript) for phishing, scam, bank fraud, OTP theft, fake prize/lottery scams, Vodafone Cash fees, or malicious links.

CRITICAL CLASSIFICATION RULES:
- Normal everyday conversations, greetings, single words like 'بريد' or 'مرحبا', test messages, or casual chats are NOT fraud (`is_fraud: false`).
- ONLY classify as fraud (`is_fraud: true`) if the message contains explicit malicious scam intent: fake prize wins requesting money/Vodafone Cash/cards/fees, urgent bank account suspension threat, request for OTP/verification code, or phishing links.

Sender: {req.sender}
Message: "{req.message_body}"

Return strictly valid JSON only without markdown formatting:
{{
  "is_fraud": true/false,
  "confidence_score": 0.98,
  "details": "سبب الخطورة بالتفصيل بالعربية (مثال: تحذير: الرسالة تحوي إغراء بجائزة وهمية وطلب تحويل أموال عبر فودافون كاش)"
}}"""

    try:
        import httpx
        import json
        async with httpx.AsyncClient(timeout=10.0) as client:
            resp = await client.post(
                "http://localhost:11434/api/generate",
                json={
                    "model": "llama3",
                    "prompt": prompt,
                    "stream": False
                }
            )
            if resp.status_code == 200:
                data = resp.json()
                raw = data.get("response", "").strip()

                clean_json = raw
                if "```json" in clean_json:
                    clean_json = clean_json.split("```json")[1].split("```")[0].strip()
                elif "```" in clean_json:
                    clean_json = clean_json.split("```")[1].split("```")[0].strip()

                parsed = json.loads(clean_json)
                return parsed
    except Exception as e:
        print(f"Ollama API notice: {e}")

    # Smart compound fallback if Ollama is offline
    msg_lower = req.message_body.lower()

    # 1. مؤشرات الجوائز والولاء وفودافون كاش والمحافظ
    is_prize_scam = ("مبروك" in msg_lower or "فزت" in msg_lower or "جائزة" in msg_lower) and ("كاش" in msg_lower or "رسوم" in msg_lower or "محفظتك" in msg_lower or "كارت" in msg_lower or "تحويل" in msg_lower or "رمز" in msg_lower or "otp" in msg_lower)

    # 2. مؤشرات البنوك والـ OTP
    is_bank_scam = ("حساب" in msg_lower or "بنك" in msg_lower or "فيزا" in msg_lower or "otp" in msg_lower or "رمز" in msg_lower) and ("إيقاف" in msg_lower or "تجميد" in msg_lower or "تحديث" in msg_lower or "بيانات" in msg_lower)

    # 3. مؤشرات الشحنات الوهمية
    is_shipping_scam = ("شحنة" in msg_lower or "طرد" in msg_lower or "البريد" in msg_lower) and ("رسوم" in msg_lower or "سدد" in msg_lower or "رابط" in msg_lower or "http" in msg_lower)

    is_fraud = is_prize_scam or is_bank_scam or is_shipping_scam

    if is_fraud:
        details = "تحذير: تم رصد محاولة احتيال (طلب تحويل أموال/رسوم إدارية عبر كاش أو سرقة بيانات بنكية ورمز OTP)"
    else:
        details = "الرسالة آمنة ومحادثة طبيعية"

    return {
        "is_fraud": is_fraud,
        "confidence_score": 0.95 if is_fraud else 0.98,
        "details": details
    }
